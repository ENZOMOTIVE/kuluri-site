from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

GREEN=RGBColor(0x17,0x3C,0x2B); CREAM=RGBColor(0xF5,0xEE,0xDF); TERRA=RGBColor(0xC6,0x5D,0x3B)
WARM=RGBColor(0xE4,0xDC,0xCD); OFFW=RGBColor(0xFF,0xFA,0xF0)
DISP="DM Serif Display"; HEAD="Montserrat"; BODY="Arimo"

prs=Presentation(); prs.slide_width=Emu(18288000); prs.slide_height=Emu(10287000)
blank=prs.slide_layouts[6]

def slide():
    s=prs.slides.add_slide(blank)
    bg=s.shapes.add_shape(1, 0,0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb=CREAM; bg.line.fill.background(); bg.shadow.inherit=False
    st=bg._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: bg._element.remove(st)
    return s

def box(s,x,y,w,h,fill=None,line=None):
    sh=s.shapes.add_shape(1, Inches(x),Inches(y),Inches(w),Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    else: sh.fill.background()
    if line: sh.line.color.rgb=line; sh.line.width=Pt(1.5)
    else: sh.line.fill.background()
    sh.shadow.inherit=False
    st=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: sh._element.remove(st)
    return sh

def tb(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=t.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    return tf

def para(tf,text,size,color,font=BODY,bold=False,space_before=0,space_after=6,first=False,line=1.15):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.text=text
    p.space_before=Pt(space_before); p.space_after=Pt(space_after); p.line_spacing=line
    for r in p.runs:
        r.font.size=Pt(size); r.font.color.rgb=color; r.font.name=font; r.font.bold=bold
    return p

# ─────────────────────────── SLIDE A — the data ───────────────────────────
s=slide()
tf=tb(s,1.0,0.75,17.0,1.1)
para(tf,"What the data told us",60,GREEN,DISP,first=True,space_after=0)
tf=tb(s,1.0,1.85,17.0,0.5)
para(tf,"Demo on your dataset farm in Milano. Our story is the Cretan olive market we are building for.",21,TERRA,BODY,first=True,space_after=0)

# left column
box(s,1.0,2.75,5.5,7.35,fill=OFFW)
tf=tb(s,1.4,3.1,4.9,6.9)
para(tf,"WHAT WE USED",17,TERRA,HEAD,bold=True,first=True,space_after=14)
para(tf,"DG AGRI synthetic dataset",20,GREEN,HEAD,bold=True,space_after=4)
para(tf,"Farm FARM-13079 · Milano, NUTS3 ITC4C\n371 ha · 62 parcels · organic · Natura 2000\nCommon wheat 137 ha · barley 83 · maize 62\nsunflower 61 · rapeseed 28\n10 years of harvest, inputs, irrigation,\nlabour, sales and certificates",17,GREEN,BODY,space_after=16)
para(tf,"Eurostat GISCO",20,GREEN,HEAD,bold=True,space_after=4)
para(tf,"NUTS 2024 outlines — the 60 km\nseason-planner map",17,GREEN,BODY,space_after=16)
para(tf,"Simulated for the demo",20,GREEN,HEAD,bold=True,space_after=4)
para(tf,"Job posts, worker pool, chats, agreements",17,GREEN,BODY,space_after=18)
para(tf,"The dataset itself is not in our repository\n(about 12 GB, git-ignored). The extracted\nfarm files are, so the app runs without it.",16,TERRA,BODY,space_after=0)

# right column — three findings
findings=[
 ("9,517 hours. Not one seasonal worker.",
  "All recorded labour is holder, family or paid-regular — on a 371 ha farm that harvests 2,600 t a year.",
  "So Kuluri records the work nobody logs: the worker confirms their own week."),
 ("20 of 30 sales cannot prove the claim.",
  "The organic certificate is valid 2024-2026 only. Two thirds of sales carry no date-valid holding certificate.",
  "So Buyer requests answers from confirmed records — totals only."),
 ("Proof has to travel to 22 buyers.",
  "22 distinct buyers over ten years, on forward, cooperative and spot contracts. 15 of 30 sales below reference price.",
  "So one record is re-used for every buyer and every CAP social-conditionality check."),
]
y=2.75
for i,(head,detail,decision) in enumerate(findings,1):
    box(s,6.9,y,11.4,2.3,fill=OFFW)
    nb=box(s,6.9,y,0.14,2.3,fill=TERRA)
    tf=tb(s,7.35,y+0.28,10.7,1.8)
    para(tf,f"{i} · {head}",24,GREEN,HEAD,bold=True,first=True,space_after=6)
    para(tf,detail,17,GREEN,BODY,space_after=7)
    para(tf,"→  "+decision,18,TERRA,BODY,bold=True,space_after=0)
    y+=2.53

tf=tb(s,1.0,10.35,17.3,0.5)
para(tf,"Synthetic dataset only · no real personal data · every figure reproducible with app/scripts/extract_farm.py · simulated demo content is labelled in the app",15,GREEN,BODY,first=True,space_after=0)

# ─────────────────────────── SLIDE B — AI use ───────────────────────────
s=slide()
tf=tb(s,1.0,0.75,17.0,1.1)
para(tf,"How we used AI",60,GREEN,DISP,first=True,space_after=0)
tf=tb(s,1.0,1.85,17.0,0.5)
para(tf,"Declared, as the Hack Guide asks. Allowed, optional, transparent.",21,TERRA,BODY,first=True,space_after=0)

# band 1
box(s,1.0,2.6,17.3,3.0,fill=OFFW)
tf=tb(s,1.5,2.95,16.3,2.6)
para(tf,"TO BUILD IT",17,TERRA,HEAD,bold=True,first=True,space_after=10)
rows=[("Scoping","Idea and dataset exploration, feasibility, market model"),
      ("App","Supervised code generation (Vue 3 / Quasar / Capacitor)"),
      ("Docs & diagrams","README, terms, citation metadata, architecture diagrams"),
      ("This deck","Structure, wording and slide generation")]
for k,v in rows:
    p=tf.add_paragraph(); p.space_after=Pt(3); p.line_spacing=1.1
    r1=p.add_run(); r1.text=f"{k}   "; r1.font.size=Pt(18); r1.font.bold=True; r1.font.name=HEAD; r1.font.color.rgb=GREEN
    r2=p.add_run(); r2.text=v; r2.font.size=Pt(18); r2.font.name=BODY; r2.font.color.rgb=GREEN
para(tf,"Every output was reviewed before it was merged and is owned by a named team member. The git history shows which commits, line by line.",17,TERRA,BODY,space_before=8,space_after=0)

# band 2 — dominant
box(s,1.0,5.85,17.3,2.5,fill=GREEN)
tf=tb(s,1.5,6.15,16.3,2.0)
para(tf,"INSIDE THE PRODUCT: NONE",17,RGBColor(0xC6,0x5D,0x3B),HEAD,bold=True,first=True,space_after=10)
para(tf,"No AI model is queried anywhere in the prototype. The chat analysis that drafts agreement terms is transparent rule-based extraction — auditable line by line, in app/src/data/agreements.js.",23,CREAM,BODY,space_after=8)
para(tf,"Deliberate: a product whose job is to make labour records trustworthy should not start with a non-deterministic core.",19,RGBColor(0xEA,0xD8,0xC8),BODY,space_after=0)

# band 3
box(s,1.0,8.55,17.3,2.15,fill=WARM)
tf=tb(s,1.5,8.8,16.3,1.75)
para(tf,"GENERATED DEMO CONTENT, AND WHAT COMES NEXT",17,TERRA,HEAD,bold=True,first=True,space_after=8)
para(tf,"AI-generated, illustrative, not verified data: the seeded chats, the worker pool, the job posts and the Romanian interface strings (machine-translated, not professionally reviewed).",16,GREEN,BODY,space_after=6)
para(tf,"Planned, with safeguards: voice transcription and translation, original always shown beside it. Any future LLM agreement drafting only with both parties' consent, always citing the messages it used.",16,GREEN,BODY,space_after=0)

prs.save("slides_data_and_ai.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
